"""Generate CareerLens AI PPTX and PDF.

The PPTX is built with python-pptx as requested. The PDF is generated with the
same layout rules using ReportLab so it is available even without PowerPoint or
LibreOffice installed.
"""

from __future__ import annotations

import sys
from pathlib import Path


PROJECT_ROOT = Path(__file__).resolve().parents[1]
LOCAL_DEPS = PROJECT_ROOT / "presentation_deps"
if LOCAL_DEPS.exists():
    sys.path.insert(0, str(LOCAL_DEPS))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas


OUTPUT_DIR = PROJECT_ROOT / "presentation_output"
PPTX_PATH = OUTPUT_DIR / "CareerLens_AI_Presentation.pptx"
PDF_PATH = OUTPUT_DIR / "CareerLens_AI_Presentation.pdf"

SLIDE_W = 13.333
SLIDE_H = 7.5
BLUE = RGBColor(43, 91, 191)
PURPLE = RGBColor(112, 72, 170)
INK = RGBColor(28, 31, 38)
MUTED = RGBColor(92, 101, 120)
LIGHT_BLUE = RGBColor(236, 242, 255)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)


def add_textbox(slide, text, x, y, w, h, size=24, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    run = paragraph.runs[0]
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_heading(slide, title):
    add_textbox(slide, title, 0.75, 0.35, 11.8, 0.55, size=30, bold=True, color=BLUE)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(1.05), Inches(1.35), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = PURPLE
    line.line.fill.background()


def add_bullets(slide, items, x=0.9, y=1.45, w=6.2, h=4.8, size=20):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Calibri"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = INK
        paragraph.space_after = Pt(9)
    return box


def add_note_block(slide, text, x, y, w, h):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT_BLUE
    shape.line.color.rgb = BLUE
    shape.line.width = Pt(1)
    frame = shape.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Calibri"
    paragraph.font.size = Pt(17)
    paragraph.font.color.rgb = INK
    paragraph.alignment = PP_ALIGN.LEFT
    return shape


def add_footer(slide, number):
    add_textbox(slide, f"CareerLens AI | {number}", 10.9, 7.08, 1.7, 0.25, size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def add_placeholder(slide, text, x, y, w, h):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = RGBColor(130, 139, 158)
    shape.line.width = Pt(1.2)
    frame = shape.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.name = "Calibri"
    paragraph.font.size = Pt(16)
    paragraph.font.color.rgb = MUTED
    return shape


def add_simple_workflow(slide):
    labels = [
        "Resume Upload",
        "Resume Parsing",
        "Text Cleaning",
        "TF-IDF Feature Extraction",
        "ML Model Prediction",
        "Top 3 Job Roles",
        "AI Description + Recommendations",
        "Final Output",
    ]
    x = 4.25
    y = 1.25
    w = 4.8
    h = 0.42
    gap = 0.26
    for index, label in enumerate(labels):
        box_y = y + index * (h + gap)
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(box_y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = BLACK
        shape.line.width = Pt(1)
        frame = shape.text_frame
        frame.clear()
        paragraph = frame.paragraphs[0]
        paragraph.text = label
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.font.name = "Calibri"
        paragraph.font.size = Pt(15)
        paragraph.font.color.rgb = BLACK
        if index < len(labels) - 1:
            cx = x + w / 2
            y1 = box_y + h
            y2 = box_y + h + gap
            line = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(cx),
                Inches(y1),
                Inches(cx),
                Inches(y2),
            )
            line.line.color.rgb = BLACK
            line.line.width = Pt(1)
            line.line.end_arrowhead = True


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    blank = prs.slide_layouts[6]

    # Slide 1
    slide = prs.slides.add_slide(blank)
    add_textbox(slide, "CareerLens AI", 0.85, 1.15, 7.5, 0.65, size=26, bold=True, color=PURPLE)
    add_textbox(slide, "Smart Resume Intelligence System", 0.85, 1.85, 9.8, 0.8, size=38, bold=True, color=BLUE)
    add_textbox(slide, "AI-Powered Resume Analysis & Career Guidance", 0.85, 2.75, 8.5, 0.4, size=20, color=MUTED)
    add_note_block(slide, "Name: Priti Darshini Biswal\nCourse: B.Tech CSE\nCollege: Gurugram University", 0.85, 4.45, 5.4, 1.25)
    add_footer(slide, 1)

    # Slide 2
    slide = prs.slides.add_slide(blank)
    add_heading(slide, "The Problem with Traditional Resumes")
    add_bullets(
        slide,
        [
            "Users don’t know suitable job roles",
            "Lack of skill gap awareness",
            "Resumes are not ATS-friendly",
            "No intelligent guidance system",
        ],
    )
    add_note_block(slide, "Example: A student with Python, SQL, and project experience may not know whether to target Data Analyst, Python Developer, or Testing roles.", 7.45, 1.55, 4.7, 2.1)
    add_footer(slide, 2)

    # Slide 3
    slide = prs.slides.add_slide(blank)
    add_heading(slide, "Introducing CareerLens AI")
    add_bullets(
        slide,
        [
            "Upload Resume (PDF)",
            "Predict Top 3 Job Roles",
            "Generate AI-based description",
            "Skill gap analysis",
            "Resume scoring and recommendations",
        ],
        x=0.9,
        y=1.55,
        w=7.0,
    )
    add_placeholder(slide, "Insert Screenshot: Application Home", 8.0, 1.65, 4.1, 3.5)
    add_footer(slide, 3)

    # Slide 4
    slide = prs.slides.add_slide(blank)
    add_heading(slide, "System Workflow")
    add_simple_workflow(slide)
    add_footer(slide, 4)

    # Slide 5
    slide = prs.slides.add_slide(blank)
    add_heading(slide, "Core ML Model")
    add_bullets(
        slide,
        [
            "Algorithm: Support Vector Machine (SVC)",
            "Feature Extraction: TF-IDF",
            "Dataset: Resume dataset",
            "Output: Top 3 predicted job roles",
        ],
        x=0.9,
        y=1.5,
        w=6.4,
    )
    add_note_block(slide, "Input → Resume text\nOutput → Data Scientist, Analyst, Python Developer", 7.35, 1.7, 4.7, 1.7)
    add_footer(slide, 5)

    # Slide 6
    slide = prs.slides.add_slide(blank)
    add_heading(slide, "AI-Powered Intelligence")
    add_bullets(
        slide,
        [
            "Resume description generation",
            "Chatbot-based interaction",
            "Smart career insights",
        ],
        x=0.9,
        y=1.55,
        w=6.2,
    )
    add_placeholder(slide, "Insert Screenshot: Chatbot Interaction", 7.5, 1.65, 4.6, 3.4)
    add_footer(slide, 6)

    # Slide 7
    slide = prs.slides.add_slide(blank)
    add_heading(slide, "System Capabilities")
    add_bullets(
        slide,
        [
            "Resume parsing from PDF",
            "Job role prediction (Top 3)",
            "Skill gap analysis",
            "Resume scoring",
            "Course recommendations",
            "Resume improvement suggestions",
        ],
        x=0.9,
        y=1.45,
        w=7.0,
    )
    add_footer(slide, 7)

    # Slide 8
    slide = prs.slides.add_slide(blank)
    add_heading(slide, "Project Implementation")
    add_placeholder(slide, "Insert Screenshot:\nResume Upload Page", 0.9, 1.55, 3.65, 3.7)
    add_placeholder(slide, "Insert Screenshot:\nPrediction Output", 4.85, 1.55, 3.65, 3.7)
    add_placeholder(slide, "Insert Screenshot:\nAI Description Output", 8.8, 1.55, 3.65, 3.7)
    add_footer(slide, 8)

    # Slide 9
    slide = prs.slides.add_slide(blank)
    add_heading(slide, "Future Scope")
    add_bullets(
        slide,
        [
            "Job matching system",
            "Resume ranking",
            "Advanced chatbot",
            "Full web deployment",
        ],
        x=0.9,
        y=1.55,
        w=6.5,
    )
    add_textbox(slide, "Thank You", 7.5, 3.0, 4.2, 0.75, size=38, bold=True, color=PURPLE, align=PP_ALIGN.CENTER)
    add_footer(slide, 9)

    OUTPUT_DIR.mkdir(exist_ok=True)
    prs.save(PPTX_PATH)


def pdf_text(c, text, x, y, size=18, color=colors.HexColor("#1C1F26"), bold=False):
    c.setFillColor(color)
    c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
    c.drawString(x, y, text)


def wrap_text(c, text, x, y, w, size=16, leading=22, color=colors.HexColor("#1C1F26"), bold=False):
    c.setFillColor(color)
    c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
    words = text.split()
    line = ""
    for word in words:
        test = f"{line} {word}".strip()
        if c.stringWidth(test, "Helvetica-Bold" if bold else "Helvetica", size) <= w:
            line = test
        else:
            c.drawString(x, y, line)
            y -= leading
            line = word
    if line:
        c.drawString(x, y, line)
    return y - leading


def pdf_heading(c, title, slide_no):
    pdf_text(c, title, 0.75 * inch, 6.78 * inch, 30, colors.HexColor("#2B5BBF"), True)
    c.setFillColor(colors.HexColor("#7048AA"))
    c.rect(0.75 * inch, 6.42 * inch, 1.35 * inch, 0.04 * inch, fill=1, stroke=0)
    pdf_text(c, f"CareerLens AI | {slide_no}", 10.9 * inch, 0.28 * inch, 9, colors.HexColor("#5C6578"))


def pdf_bullets(c, items, x=0.9 * inch, y=5.85 * inch, size=20):
    c.setFont("Helvetica", size)
    c.setFillColor(colors.HexColor("#1C1F26"))
    for item in items:
        c.circle(x, y + 5, 2.4, fill=1, stroke=0)
        wrap_text(c, item, x + 0.22 * inch, y, 6.2 * inch, size=size, leading=0.34 * inch)
        y -= 0.55 * inch


def pdf_note(c, text, x, y, w, h, size=16):
    c.setFillColor(colors.HexColor("#ECF2FF"))
    c.setStrokeColor(colors.HexColor("#2B5BBF"))
    c.rect(x, y, w, h, fill=1, stroke=1)
    ty = y + h - 0.32 * inch
    for line in text.split("\n"):
        ty = wrap_text(c, line, x + 0.18 * inch, ty, w - 0.36 * inch, size=size, leading=0.25 * inch)


def pdf_placeholder(c, text, x, y, w, h):
    c.setFillColor(colors.white)
    c.setStrokeColor(colors.HexColor("#828B9E"))
    c.rect(x, y, w, h, fill=1, stroke=1)
    c.setFillColor(colors.HexColor("#5C6578"))
    c.setFont("Helvetica", 16)
    lines = text.split("\n")
    start_y = y + h / 2 + (len(lines) - 1) * 10
    for idx, line in enumerate(lines):
        c.drawCentredString(x + w / 2, start_y - idx * 22, line)


def pdf_workflow(c):
    labels = [
        "Resume Upload",
        "Resume Parsing",
        "Text Cleaning",
        "TF-IDF Feature Extraction",
        "ML Model Prediction",
        "Top 3 Job Roles",
        "AI Description + Recommendations",
        "Final Output",
    ]
    x = 4.25 * inch
    y_top = 5.83 * inch
    w = 4.8 * inch
    h = 0.42 * inch
    gap = 0.26 * inch
    for index, label in enumerate(labels):
        y = y_top - index * (h + gap)
        c.setFillColor(colors.white)
        c.setStrokeColor(colors.black)
        c.rect(x, y, w, h, fill=1, stroke=1)
        c.setFillColor(colors.black)
        c.setFont("Helvetica", 15)
        c.drawCentredString(x + w / 2, y + 0.15 * inch, label)
        if index < len(labels) - 1:
            cx = x + w / 2
            c.line(cx, y, cx, y - gap + 0.05 * inch)
            c.line(cx, y - gap + 0.05 * inch, cx - 0.05 * inch, y - gap + 0.13 * inch)
            c.line(cx, y - gap + 0.05 * inch, cx + 0.05 * inch, y - gap + 0.13 * inch)


def build_pdf():
    pagesize = landscape((SLIDE_H * inch, SLIDE_W * inch))
    c = canvas.Canvas(str(PDF_PATH), pagesize=pagesize)

    # 1
    pdf_text(c, "CareerLens AI", 0.85 * inch, 6.0 * inch, 26, colors.HexColor("#7048AA"), True)
    pdf_text(c, "Smart Resume Intelligence System", 0.85 * inch, 5.35 * inch, 38, colors.HexColor("#2B5BBF"), True)
    pdf_text(c, "AI-Powered Resume Analysis & Career Guidance", 0.85 * inch, 4.9 * inch, 20, colors.HexColor("#5C6578"))
    pdf_note(c, "Name: Priti Darshini Biswal\nCourse: B.Tech CSE\nCollege: Gurugram University", 0.85 * inch, 1.8 * inch, 5.4 * inch, 1.25 * inch, 17)
    pdf_text(c, "CareerLens AI | 1", 10.9 * inch, 0.28 * inch, 9, colors.HexColor("#5C6578"))
    c.showPage()

    pdf_heading(c, "The Problem with Traditional Resumes", 2)
    pdf_bullets(c, ["Users don’t know suitable job roles", "Lack of skill gap awareness", "Resumes are not ATS-friendly", "No intelligent guidance system"])
    pdf_note(c, "Example: A student with Python, SQL, and project experience may not know whether to target Data Analyst, Python Developer, or Testing roles.", 7.45 * inch, 3.75 * inch, 4.7 * inch, 2.1 * inch, 17)
    c.showPage()

    pdf_heading(c, "Introducing CareerLens AI", 3)
    pdf_bullets(c, ["Upload Resume (PDF)", "Predict Top 3 Job Roles", "Generate AI-based description", "Skill gap analysis", "Resume scoring and recommendations"])
    pdf_placeholder(c, "Insert Screenshot:\nApplication Home", 8.0 * inch, 2.35 * inch, 4.1 * inch, 3.5 * inch)
    c.showPage()

    pdf_heading(c, "System Workflow", 4)
    pdf_workflow(c)
    c.showPage()

    pdf_heading(c, "Core ML Model", 5)
    pdf_bullets(c, ["Algorithm: Support Vector Machine (SVC)", "Feature Extraction: TF-IDF", "Dataset: Resume dataset", "Output: Top 3 predicted job roles"])
    pdf_note(c, "Input → Resume text\nOutput → Data Scientist, Analyst, Python Developer", 7.35 * inch, 4.1 * inch, 4.7 * inch, 1.7 * inch, 17)
    c.showPage()

    pdf_heading(c, "AI-Powered Intelligence", 6)
    pdf_bullets(c, ["Resume description generation", "Chatbot-based interaction", "Smart career insights"])
    pdf_placeholder(c, "Insert Screenshot:\nChatbot Interaction", 7.5 * inch, 2.45 * inch, 4.6 * inch, 3.4 * inch)
    c.showPage()

    pdf_heading(c, "System Capabilities", 7)
    pdf_bullets(c, ["Resume parsing from PDF", "Job role prediction (Top 3)", "Skill gap analysis", "Resume scoring", "Course recommendations", "Resume improvement suggestions"])
    c.showPage()

    pdf_heading(c, "Project Implementation", 8)
    pdf_placeholder(c, "Insert Screenshot:\nResume Upload Page", 0.9 * inch, 2.25 * inch, 3.65 * inch, 3.7 * inch)
    pdf_placeholder(c, "Insert Screenshot:\nPrediction Output", 4.85 * inch, 2.25 * inch, 3.65 * inch, 3.7 * inch)
    pdf_placeholder(c, "Insert Screenshot:\nAI Description Output", 8.8 * inch, 2.25 * inch, 3.65 * inch, 3.7 * inch)
    c.showPage()

    pdf_heading(c, "Future Scope", 9)
    pdf_bullets(c, ["Job matching system", "Resume ranking", "Advanced chatbot", "Full web deployment"])
    pdf_text(c, "Thank You", 7.8 * inch, 3.7 * inch, 38, colors.HexColor("#7048AA"), True)
    c.showPage()

    c.save()


def main():
    OUTPUT_DIR.mkdir(exist_ok=True)
    build_pptx()
    build_pdf()
    print(PPTX_PATH)
    print(PDF_PATH)


if __name__ == "__main__":
    main()
